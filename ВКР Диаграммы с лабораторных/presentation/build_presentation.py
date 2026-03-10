from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path("/Users/glebporosin/AndroidStudioProjects/ECommerceARapp/ВКР Диаграммы с лабораторных/presentation")
ASSETS_DIR = BASE_DIR / "assets"
OUTPUT_PPTX = BASE_DIR / "architecture-presentation.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(245, 247, 250)
PANEL = RGBColor(255, 255, 255)
NAVY = RGBColor(13, 58, 107)
BLUE = RGBColor(16, 97, 176)
TEAL = RGBColor(35, 162, 217)
GREEN = RGBColor(130, 179, 102)
ORANGE = RGBColor(215, 155, 0)
TEXT = RGBColor(32, 37, 43)
MUTED = RGBColor(97, 105, 115)


def add_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    top_band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.42)
    )
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = NAVY
    top_band.line.fill.background()


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.52), Inches(8.6), Inches(0.7))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Times New Roman"
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = NAVY
    if subtitle:
        p = tf.add_paragraph()
        p.space_before = Pt(3)
        run = p.add_run()
        run.text = subtitle
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED


def add_footer(slide, text):
    box = slide.shapes.add_textbox(Inches(10.35), Inches(7.05), Inches(2.5), Inches(0.25))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def add_bullets_box(slide, x, y, w, h, heading, bullets, accent=BLUE):
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = accent
    panel.line.width = Pt(1.5)

    header = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, Inches(0.43)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = accent
    header.line.fill.background()

    header_tf = header.text_frame
    header_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = header_tf.paragraphs[0]
    run = p.add_run()
    run.text = heading
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.52), w - Inches(0.3), h - Inches(0.62))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {bullet}"
        p.level = 0
        p.space_after = Pt(6)
        p.font.name = "Times New Roman"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT


def add_caption(slide, x, y, w, h, text, accent=GREEN):
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = accent
    panel.line.width = Pt(1.5)
    tf = panel.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(10)
    tf.margin_right = Pt(10)
    tf.margin_top = Pt(8)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = TEXT


def add_image(slide, image_path, x, y, w, h):
    slide.shapes.add_picture(str(image_path), x, y, width=w, height=h)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)

    title = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(7.9), Inches(1.6))
    tf = title.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Архитектура мобильного\nприложения магазина мебели\nс AR-примеркой"
    run.font.name = "Times New Roman"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = NAVY

    roadmap = [
        "C4 Context: кто использует систему и с какими внешними системами она взаимодействует",
        "C4 Container: из каких технических контейнеров состоит решение",
        "Domain Model: какие бизнес-сущности и value objects лежат в основе предметной области",
        "UML Core + Scaled Architecture: как реализован ключевой PDP-AR-сценарий и как ядро связано с инфраструктурой",
        "Facade Diagram: как скрыта сложность iOS AR-подсистемы",
    ]
    add_bullets_box(slide, Inches(8.55), Inches(1.1), Inches(4.1), Inches(5.4), "Что внутри", roadmap, accent=BLUE)
    add_footer(slide, "Подготовлено автоматически из draw.io диаграмм")


def image_plus_text_slide(prs, title, subtitle, image_name, heading, bullets, footer, accent=BLUE, caption=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, title, subtitle)

    img_panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.35), Inches(8.25), Inches(5.65)
    )
    img_panel.fill.solid()
    img_panel.fill.fore_color.rgb = PANEL
    img_panel.line.color.rgb = RGBColor(215, 220, 227)
    img_panel.line.width = Pt(1.0)

    add_image(slide, ASSETS_DIR / image_name, Inches(0.72), Inches(1.52), Inches(7.92), Inches(5.28))

    add_bullets_box(slide, Inches(9.0), Inches(1.35), Inches(3.78), Inches(4.9), heading, bullets, accent=accent)
    if caption:
        add_caption(slide, Inches(9.0), Inches(6.1), Inches(3.78), Inches(0.9), caption, accent=accent)
    add_footer(slide, footer)


def conclusion_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Итог", "Логика проектирования от бизнес-уровня к реализации")

    bullets = [
        "Контекстная и контейнерная диаграммы задают границы системы и показывают техническое разбиение решения.",
        "Доменная модель фиксирует устойчивые бизнес-сущности: товар, корзину и AR-описание продукта.",
        "Core-диаграмма показывает ключевой пользовательский сценарий PDP -> кэш модели -> AR -> корзина.",
        "Scaled Architecture объясняет, как UI, use cases, repository contracts и adapters связаны по правилам Clean Architecture.",
        "Facade Diagram демонстрирует локальное упрощение сложной iOS AR-подсистемы через единый ARSceneFacade.",
    ]
    add_bullets_box(slide, Inches(0.7), Inches(1.45), Inches(7.2), Inches(4.95), "Главные выводы", bullets, accent=GREEN)
    add_caption(
        slide,
        Inches(8.2),
        Inches(1.45),
        Inches(4.45),
        Inches(2.2),
        "Вся серия диаграмм описывает один и тот же проект на разных уровнях абстракции: сначала внешнее окружение, затем внутреннюю структуру, потом доменную модель, ключевой сценарий и частный архитектурный паттерн.",
        accent=ORANGE,
    )
    add_caption(
        slide,
        Inches(8.2),
        Inches(3.95),
        Inches(4.45),
        Inches(1.55),
        "Такой набор схем можно использовать и в презентации, и в пояснительной записке: каждая диаграмма отвечает на свой архитектурный вопрос и не дублирует предыдущую.",
        accent=BLUE,
    )
    add_footer(slide, "Финальный слайд")


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title_slide(prs)

    image_plus_text_slide(
        prs,
        "C4 Context",
        "Кто взаимодействует с системой и какие внешние зависимости есть у проекта",
        "lab1-context.png",
        "Что показывает диаграмма",
        [
            "Система рассматривается как единый программный продукт магазина мебели с AR-примеркой.",
            "Основной актор: покупатель, который просматривает каталог, открывает карточку товара, запускает AR и управляет корзиной.",
            "Внешние системы: CDN/Object Storage для медиа и 3D-моделей, а также ARCore/ARKit как платформенные AR-сервисы устройства.",
        ],
        "Лабораторная 1, страница 1",
        accent=BLUE,
        caption="Эта схема нужна, чтобы зафиксировать границы системы и ключевые пользовательские use cases до обсуждения технологий.",
    )

    image_plus_text_slide(
        prs,
        "C4 Container",
        "Из каких технических контейнеров состоит решение и где живут данные",
        "lab1-container.png",
        "Что показывает диаграмма",
        [
            "Пользователь работает через мобильное приложение, которое получает данные каталога через Backend API.",
            "Карточки товаров и ссылки на медиа живут в каталоге, а корзина и metadata моделей хранятся локально.",
            "AR-сценарий опирается на локальный файловый кэш 3D-моделей и на платформенные сервисы ARCore/ARKit.",
        ],
        "Лабораторная 1, страница 2",
        accent=TEAL,
        caption="Эта схема раскрывает систему до уровня контейнеров и показывает, как пользовательские сценарии распределяются по техническим блокам.",
    )

    image_plus_text_slide(
        prs,
        "Domain Model",
        "Какие бизнес-сущности лежат в основе предметной области проекта",
        "lab1-domain.png",
        "Что показывает диаграмма",
        [
            "Центральные сущности: Product и ShoppingCart; CartLine живет внутри корзины как часть агрегата.",
            "Value Objects выделены отдельно: Money, ProductImage, ProductAttribute, ArModel, ProductDimensions, CartItemSnapshot.",
            "Смысл диаграммы не в БД, а в бизнес-логике: поведении сущностей, связях и инвариантах предметной области.",
        ],
        "Лабораторная 1, страница 3",
        accent=GREEN,
        caption="Здесь фиксируется доменная модель без SQL-типов и таблиц: сначала сущности и их смысл, потом уже возможный ORM-маппинг.",
    )

    image_plus_text_slide(
        prs,
        "UML Core Domain: PDP AR Flow",
        "Ключевой сценарий внутри ядра приложения",
        "lab2-core.png",
        "Что показывает диаграмма",
        [
            "PdpViewModel оркестрирует сценарий карточки товара и знает только use case-контракты.",
            "Сценарий: загрузить ProductPageInfo, проверить локальный кэш модели, при необходимости скачать 3D-файл и выпустить OpenArObject.",
            "Параллельно через отдельные use cases поддерживается наблюдение и изменение количества товара в локальной корзине.",
        ],
        "Лабораторная 2, страница 1",
        accent=ORANGE,
        caption="Эта схема показывает core-логику ключевого AR-сценария без привязки к конкретной инфраструктуре.",
    )

    image_plus_text_slide(
        prs,
        "Scaled Architecture",
        "Как UI, core и инфраструктура соединяются в общем архитектурном контуре",
        "lab2-scaled.png",
        "Что показывает диаграмма",
        [
            "Слева расположен UI, в центре core со ViewModel, use cases и repository contracts, справа adapters и внешние системы.",
            "UI не ходит напрямую в сеть, БД или файловую систему: все идет через ViewModel и use case-контракты.",
            "Диаграмма иллюстрирует Clean Architecture, DIP и возможность менять concrete-реализации без переписывания ядра.",
        ],
        "Лабораторная 2, страница 2",
        accent=BLUE,
        caption="Главная мысль: зависимости направлены к абстракциям, а инфраструктура подключается как внешний слой через adapters.",
    )

    image_plus_text_slide(
        prs,
        "Facade Pattern: iOS AR Subsystem",
        "Как скрыта сложность AR-подсистемы на iOS",
        "lab3-facade.png",
        "Что показывает диаграмма",
        [
            "Клиентом выступает ArScreen, который работает не с ARKit напрямую, а с единой точкой входа ARSceneFacade.",
            "Facade координирует внутренние сервисы: конфигурацию сессии, загрузку модели, размещение объекта, жесты, overlay и telemetry.",
            "Внешняя сложность RealityKit/ARKit и локального model cache скрыта за коротким бизнес-интерфейсом фасада.",
        ],
        "Лабораторная 3, страница 1",
        accent=GREEN,
        caption="Эта схема показывает применение структурного паттерна Facade для упрощения сложной AR-подсистемы.",
    )

    conclusion_slide(prs)
    prs.save(OUTPUT_PPTX)


if __name__ == "__main__":
    build()
